"""Δέουσα επιμέλεια (taxmatch/aml): μοντέλο βαθμολόγησης, αποθήκευση, ημερολόγιο, εξαγωγή."""
from datetime import date

import pytest
from openpyxl import load_workbook

from taxmatch import settings_store
from taxmatch.aml import content, export, model, store
from taxmatch.business_profiles import service
from taxmatch.deadlines import events_between

AFM = "094259216"
AFM2 = "123456783"


# ------------------------------------------------------------------ μοντέλο

def test_simple_domestic_bookkeeping_client_is_low_in_both_models():
    r = model.compute(["a_simple", "b_eu", "g_basic_small", "d_face"])
    assert r.sum_a == -2 and r.cat_a == model.LOW
    assert r.subs_b == {"A": 5, "B": 0, "G": 0, "D": 0} and r.total_b == 5 and r.cat_b == model.LOW


def test_cash_intensive_client_is_medium_not_low():
    r = model.compute(["a_simple", "a_cash", "b_eu", "g_basic", "d_face"])
    assert r.sum_a == 2 and r.cat_a == model.MEDIUM
    assert r.total_b == 15 + 0 + 5 + 0 and r.cat_b == model.MEDIUM


def test_negative_sum_with_annex_ii_factor_is_not_simplified():
    # Άθροισμα ≤ 0 αλλά με παράγοντα Π.ΙΙ (εξ αποστάσεως με eID, «εξουδετερωμένο») -> όχι απλουστευμένη (άρθρο 15 παρ. 1)
    r = model.compute(["a_public", "b_eu", "d_remote_eid"])
    assert r.sum_a == -3 and r.n_annex_ii == 1
    assert r.cat_a == model.MEDIUM
    assert any("Παραρτήματος ΙΙ" in n for n in r.notes)


def test_no_annex_i_factor_means_no_simplified_even_with_zero_sum():
    r = model.compute(["a_simple", "g_basic", "d_face"])
    assert r.sum_a == 0 and r.cat_a == model.MEDIUM


def test_heavy_factors_make_high_in_model_a():
    r = model.compute(["a_complex", "a_nominee", "b_eu"])
    assert r.sum_a == 7 and r.cat_a == model.HIGH
    assert r.subs_b["A"] == 25                        # 20 + 5 για τον δεύτερο, ταβάνι 25


def test_model_b_axis_cap_and_high_threshold():
    r = model.compute(["a_complex", "a_nominee", "a_vehicle", "b_weak", "g_anonymity", "d_third_payer"])
    assert r.subs_b == {"A": 25, "B": 20, "G": 20, "D": 20} and r.total_b == 85 and r.cat_b == model.HIGH


@pytest.mark.parametrize("override", [o.key for o in model.OVERRIDES])
def test_every_override_forces_high_even_with_negative_sum(override):
    r = model.compute(["a_listed", "b_eu", "g_basic_small"], [override])
    assert r.sum_a < 0
    assert r.cat_a == model.HIGH and r.cat_b == model.HIGH


def test_par3_transactions_add_a_note_regardless_of_category():
    r = model.compute(["a_simple", "b_eu", "d_unusual_tx"])
    assert r.par3 and r.sum_a == 4 and r.cat_a == model.MEDIUM
    assert any("παρ. 3 του άρθρου 16" in n for n in r.notes)


def test_unknown_keys_are_ignored():
    assert model.compute(["nope", "b_eu"], ["also_nope"]).sum_a == -1


def test_final_category_can_only_go_up():
    assert model.final_category(model.LOW, model.HIGH) == model.HIGH
    assert model.final_category(model.HIGH, model.LOW) == model.HIGH
    assert model.final_category(model.MEDIUM, "") == model.MEDIUM


def test_pep_profile_forces_pep_override():
    assert model.overrides_from_profile("family") == {"o_pep"}
    assert model.overrides_from_profile("no") == set()
    assert model.overrides_from_profile("unknown") == set()


def test_suggestions_from_kad_are_hints_with_reasons():
    s = model.suggest({"name": "ΤΑΒΕΡΝΑ ΑΕ", "kads": [{"code": "56.10.11.01", "descr": "Εστιατόρια"},
                                                   {"code": "68.31.11.01", "descr": "Μεσιτεία ακινήτων"}]})
    assert "a_cash" in s and "56.10.11.01" in s["a_cash"]
    assert "g_sector" in s
    assert "b_eu" in s                                    # προεπιλογή: ελληνικός ΑΦΜ


def test_every_factor_and_override_is_well_formed():
    keys = [f.key for f in model.FACTORS] + [o.key for o in model.OVERRIDES]
    assert len(keys) == len(set(keys))
    assert {f.axis for f in model.FACTORS} == {a for a, _ in model.AXES}
    assert all(f.annex in ("", "I", "II") for f in model.FACTORS)
    # Παράγοντες Π.Ι δεν προσθέτουν βαθμούς, παράγοντες Π.ΙΙ δεν αφαιρούν
    assert all(f.a <= 0 for f in model.FACTORS if f.annex == "I")
    assert all(f.a >= 0 for f in model.FACTORS if f.annex == "II")


def test_content_has_all_23_questions_and_15_steps():
    assert [q.n for q in content.QUESTIONS] == list(range(1, 24))
    assert [s.n for s in content.STEPS] == list(range(1, 16))
    assert all(0 <= q.step <= 15 for q in content.QUESTIONS)


def test_add_months_clamps_to_month_end():
    assert store.add_months(date(2026, 1, 31), 1) == date(2026, 2, 28)
    assert store.add_months(date(2026, 11, 15), 24) == date(2028, 11, 15)


# ------------------------------------------------------------------ αποθήκευση

def test_save_assessment_recomputes_and_schedules_review(conn):
    service.add(conn, AFM, "ΔΟΚΙΜΗ ΑΕ")
    aid = store.save_assessment(conn, AFM, factors=["a_simple", "b_eu", "g_basic_small", "d_face"],
                                assessed_on=date(2026, 9, 1), justification="απλός πελάτης")
    a = store.get_assessment(conn, aid)
    assert a["final_category"] == model.LOW and a["model"] == "A" and a["client_name"] == "ΔΟΚΙΜΗ ΑΕ"
    assert a["next_review"] == "2028-09-01"               # χαμηλός: 24 μήνες (προεπιλογή)


def test_review_interval_and_model_follow_office_settings(conn):
    service.add(conn, AFM, "ΔΟΚΙΜΗ ΑΕ")
    settings_store.set_value(conn, "aml_review_medium", "6")
    settings_store.set_value(conn, "aml_model", "B")
    aid = store.save_assessment(conn, AFM, factors=["a_cash", "b_eu"], assessed_on=date(2026, 9, 1))
    a = store.get_assessment(conn, aid)
    assert a["model"] == "B" and a["final_category"] == model.MEDIUM and a["next_review"] == "2027-03-01"


def test_escalation_requires_a_reason_and_is_kept(conn):
    service.add(conn, AFM, "ΔΟΚΙΜΗ ΑΕ")
    with pytest.raises(ValueError):
        store.save_assessment(conn, AFM, factors=["a_simple", "b_eu", "g_basic_small"], escalate_to=model.HIGH)
    aid = store.save_assessment(conn, AFM, factors=["a_simple", "b_eu", "g_basic_small"], escalate_to=model.HIGH,
                                escalation_note="αρνητικά δημοσιεύματα")
    a = store.get_assessment(conn, aid)
    assert a["model_category"] == model.LOW and a["final_category"] == model.HIGH
    assert a["escalation_note"] == "αρνητικά δημοσιεύματα"


def test_profile_roundtrip_and_kyc_completeness(conn):
    store.save_profile(conn, AFM, pep_status="no", purpose="Τήρηση βιβλίων, πηγή: συνέντευξη",
                       kyc={"identity": "2026-09-01", "bogus": "x", "ubo": ""})
    p = store.get_profile(conn, AFM)
    assert p["kyc"] == {"identity": "2026-09-01"}
    assert store.kyc_completeness(p) == (3, len(content.KYC_ITEMS))    # identity + purpose + pep_check


def test_overview_status_and_history_survives_client_deletion(conn):
    service.add(conn, AFM, "ΠΑΛΙΟΣ ΠΕΛΑΤΗΣ")
    service.add(conn, AFM2, "ΝΕΟΣ ΠΕΛΑΤΗΣ")
    store.save_assessment(conn, AFM, factors=["a_cash"], assessed_on=date(2025, 1, 10))   # επανεξέταση 2026-01-10
    rows = {r["afm"]: r for r in store.overview(conn, today=date(2026, 9, 23))}
    assert rows[AFM]["status"] == "overdue" and rows[AFM2]["status"] == "missing"
    counts = store.summary_counts(conn, today=date(2026, 9, 23))
    assert counts["overdue"] == 1 and counts["missing"] == 1 and counts["medium"] == 1

    service.delete(conn, AFM)
    assert store.history(conn, AFM)                        # τήρηση 5ετίας: ΔΕΝ σβήνει με τον πελάτη
    rows = {r["afm"]: r for r in store.overview(conn, today=date(2026, 9, 23))}
    assert rows[AFM]["is_client"] is False and rows[AFM]["name"] == "ΠΑΛΙΟΣ ΠΕΛΑΤΗΣ"
    assert store.summary_counts(conn, today=date(2026, 9, 23))["clients"] == 1


def test_review_dates_appear_in_calendar_only_when_asked(conn):
    service.add(conn, AFM, "ΔΟΚΙΜΗ ΑΕ")
    store.save_assessment(conn, AFM, factors=["a_cash"], assessed_on=date(2025, 10, 5))    # -> 2026-10-05
    start, end = date(2026, 10, 1), date(2026, 10, 31)
    assert not any(e["kind"] == "aml" for e in events_between(conn, start, end))
    evs = [e for e in events_between(conn, start, end, include_aml=True) if e["kind"] == "aml"]
    assert len(evs) == 1 and evs[0]["date"] == "2026-10-05" and evs[0]["afm"] == AFM
    assert [e["kind"] for e in events_between(conn, start, end, afm=AFM2, include_aml=True)].count("aml") == 0


def test_only_latest_assessment_drives_the_review_date(conn):
    service.add(conn, AFM, "ΔΟΚΙΜΗ ΑΕ")
    store.save_assessment(conn, AFM, factors=["a_cash"], assessed_on=date(2025, 10, 5))
    store.save_assessment(conn, AFM, factors=["a_cash"], assessed_on=date(2026, 9, 1), kind="periodic")
    assert store.reviews_between(conn, date(2026, 10, 1), date(2026, 10, 31)) == []
    assert store.latest(conn, AFM)["kind"] == "periodic"


def test_office_items_and_registers(conn):
    store.set_office_item(conn, "q01", "no", "να οργανωθεί σεμινάριο")
    store.set_office_item(conn, "q01", "yes")                      # χωρίς note: κρατά την προηγούμενη
    assert store.office_items(conn)["q01"] == {"state": "yes", "note": "να οργανωθεί σεμινάριο",
                                               "updated_at": store.office_items(conn)["q01"]["updated_at"]}
    rid = store.add_register(conn, "training", "2026-09-10", "Σεμινάριο AML", details="3 άτομα")
    store.add_register(conn, "report", "2026-09-11", "Ασυνήθιστη κατάθεση", afm=AFM, outcome="δεν αναφέρθηκε — αιτιολογία")
    assert [r["kind"] for r in store.list_register(conn)] == ["report", "training"]
    store.delete_register(conn, rid)
    assert [r["kind"] for r in store.list_register(conn, "training")] == []


# ------------------------------------------------------------------ εξαγωγή

def test_exports_produce_readable_workbooks(conn, tmp_path):
    service.add(conn, AFM, "ΔΟΚΙΜΗ ΑΕ")
    service.add(conn, AFM2, "ΧΩΡΙΣ ΑΞΙΟΛΟΓΗΣΗ")
    store.save_profile(conn, AFM, pep_status="domestic", purpose="Λογιστική υποστήριξη")
    aid = store.save_assessment(conn, AFM, factors=["a_simple", "b_eu"], overrides=["o_pep"],
                                justification="ΠΕΠ", approved_by="Διαχειριστής")
    p1 = export.assessment_sheet(conn, aid, tmp_path / "sheet.xlsx")
    ws = load_workbook(p1).active
    text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "ΔΟΚΙΜΗ ΑΕ" in text and "ΥΨΗΛΟΣ" in text and "Αυξημένη" in text and "ΠΕΠ ημεδαπό" in text

    p2 = export.summary_sheet(conn, tmp_path / "summary.xlsx", version="1.0")
    ws = load_workbook(p2).active
    values = [[c.value for c in row] for row in ws.iter_rows(min_row=5)]
    by_afm = {v[0]: v for v in values}
    assert by_afm[AFM][2] == "ΥΨΗΛΟΣ" and by_afm[AFM2][2] == "ΧΩΡΙΣ ΑΞΙΟΛΟΓΗΣΗ"


# ------------------------------------------------------------------ προφίλ συναλλαγών / έγγραφα / φάκελος

def test_weekly_cash_in_real_estate_is_not_low_risk():
    # Στο εργαλείο που αθροίζει πόντους ανά συναλλαγή, αυτή η γραμμή (15–50 χιλ., μετρητά 3–10 χιλ., εβδομαδιαία,
    # ακίνητα, χώρα FATF) βγαίνει 15 πόντοι = «Χαμηλός / απλοποιημένη». Εδώ γίνεται παράγοντες Παραρτήματος.
    tx = {"amount": "15_50", "channel": "cash_1_10k", "frequency": "weekly", "activity": "real_estate", "country": "fatf_ok"}
    factors, overrides, _w = model.factors_from_transactions([tx])
    assert {"a_cash", "g_sector", "b_third_ok"} <= set(factors) and not overrides
    r = model.compute(list(factors) + ["a_simple", "g_basic", "d_face"])
    assert r.cat_a == model.HIGH or r.cat_a == model.MEDIUM
    assert r.cat_a != model.LOW and r.cat_b != model.LOW


def test_transaction_country_lists_become_overrides_and_big_amounts_need_justification():
    f, o, w = model.factors_from_transactions([{"amount": "100_250", "channel": "bank", "country": "eu_high_risk"}])
    assert "o_high_risk_country" in o and "d_unusual_tx" in f and any("αιτιολόγηση" in x for x in w)
    f, o, w = model.factors_from_transactions([{"amount": "100_250", "channel": "bank", "country": "gr",
                                                "justification": "πώληση ακινήτου με συμβόλαιο"}])
    assert "d_unusual_tx" not in f and f == {"b_eu": f["b_eu"]}
    f, _o, _w = model.factors_from_transactions([], fee_payment="third")
    assert "d_third_payer" in f


def test_documents_depend_on_legal_form():
    assert "statute" in dict(content.documents_for("ΙΚΕ")) and "ubo_registry" in dict(content.documents_for("Α.Ε."))
    assert "id" in dict(content.documents_for("ΑΤΟΜΙΚΗ ΕΠΙΧΕΙΡΗΣΗ")) and "statute" not in dict(content.documents_for(""))
    assert "income" not in content.required_documents("")                     # προαιρετικό


def test_profile_keeps_transactions_docs_status_and_overview_counts(conn):
    service.add(conn, AFM, "ΔΟΚΙΜΗ ΙΚΕ")
    conn.execute("UPDATE businesses SET legal_form='ΙΚΕ' WHERE afm=?", (AFM,))
    store.save_profile(conn, AFM, pep_status="no", transactions=[{"amount": "lt15", "channel": "bank", "bogus": "x"},
                                                                  {"channel": "not-a-channel"}],
                       file_status="red_flag", docs={"decl": "2026-09-01", "nope": "x"}, ubo_state="identified",
                       fee_payment="iris")
    p = store.get_profile(conn, AFM)
    assert p["transactions"] == [{"amount": "lt15", "channel": "bank", "frequency": "", "activity": "", "country": "",
                                  "justification": ""}]
    assert p["docs"] == {"decl": "2026-09-01"} and p["file_status"] == "red_flag" and p["fee_payment"] == "iris"
    store.save_profile(conn, AFM, pep_status="no", purpose="x")                  # χωρίς τα νέα πεδία: μένουν ως έχουν
    assert store.get_profile(conn, AFM)["file_status"] == "red_flag"
    row = {r["afm"]: r for r in store.overview(conn)}[AFM]
    assert row["docs_done"] == 1 and row["docs_total"] == len(content.required_documents("ΙΚΕ"))
    counts = store.summary_counts(conn)
    assert counts["docs_missing"] == 1 and counts["alerts"] == 1


# ------------------------------------------------------------------ είδος πελάτη / υποθέσεις / έγγραφα PDF

@pytest.mark.parametrize("legal_form,kads,expected", [
    ("ΙΚΕ", 1, "company"), ("Α.Ε.", 1, "company"), ("ΕΠΕ", 1, "company"), ("Ο.Ε.", 1, "partnership"),
    ("Ε.Ε.", 1, "partnership"), ("ΑΤΟΜΙΚΗ ΕΠΙΧΕΙΡΗΣΗ", 1, "sole"), ("ΙΔΙΩΤΗΣ", 0, "individual"),
    ("", 0, "individual"), ("", 1, "sole"), ("ΝΠΔΔ", 0, "public"), ("TRUST", 0, "trust"),
])
def test_client_kind_from_legal_form(legal_form, kads, expected):
    assert content.client_kind(legal_form, kads) == expected


def test_documents_differ_by_kind_and_user_can_override_kind(conn):
    assert "partners_id" in dict(content.documents_for("", "partnership"))
    assert "start" in dict(content.documents_for("", "sole")) and "statute" not in dict(content.documents_for("", "sole"))
    assert "trust_parties" in dict(content.documents_for("", "trust"))
    service.add(conn, AFM, "ΠΕΛΑΤΗΣ")
    conn.execute("UPDATE businesses SET activity_state='none' WHERE afm=?", (AFM,))     # ΙΔΙΩΤΗΣ κατά ΑΑΔΕ
    assert {r["afm"]: r for r in store.overview(conn)}[AFM]["client_kind"] == "individual"
    store.save_profile(conn, AFM, pep_status="no", client_kind="company")
    row = {r["afm"]: r for r in store.overview(conn)}[AFM]
    assert row["client_kind"] == "company" and row["docs_total"] == len(content.required_documents("", "company"))


def test_whistle_case_deadlines_and_attention():
    from taxmatch.aml import cases
    c = {"kind": "whistle", "received_on": "2026-09-23"}
    d = cases.deadlines(c)
    assert d["ack_due"] == "2026-10-02"            # 7 εργάσιμες (Σαβ/Κυρ εξαιρούνται)
    assert d["feedback_due"] == "2026-12-23"       # 3 μήνες


def test_cases_validation_and_attention(conn):
    from taxmatch.aml import cases
    service.add(conn, AFM, "ΠΕΛΑΤΗΣ")
    with pytest.raises(ValueError):                 # ύποπτη συναλλαγή χωρίς πελάτη
        cases.save(conn, {"kind": "suspicion", "received_on": "2026-01-01", "description": "x"})
    with pytest.raises(ValueError):                 # μη αναφορά χωρίς αιτιολογία
        cases.save(conn, {"kind": "suspicion", "received_on": "2026-01-01", "description": "x", "afm": AFM,
                          "decision": "no_report"})
    sid = cases.save(conn, {"kind": "suspicion", "received_on": "2026-01-01", "description": "μετρητά",
                            "afm": AFM, "red_flags": ["structuring", "bogus"]})
    wid = cases.save(conn, {"kind": "whistle", "received_on": "2026-01-01", "description": "πίεση",
                            "anonymous": True, "reporter": "Κάποιος"})
    s = cases.get(conn, sid)
    assert s["red_flags"] == ["structuring"] and s["client_name"] == "ΠΕΛΑΤΗΣ"
    assert cases.get(conn, wid)["reporter"] == ""   # ανωνυμία: δεν αποθηκεύεται όνομα
    reasons = {c["id"]: c["reason"] for c in cases.attention(conn, today=date(2026, 9, 23))}
    assert "απόφαση" in reasons[sid] and "βεβαίωση" in reasons[wid]
    cases.save(conn, {**cases.get(conn, sid), "decision": "report", "status": "closed"}, sid)
    assert sid not in {c["id"] for c in cases.attention(conn, today=date(2026, 9, 23))}
    service.delete(conn, AFM)
    assert cases.get(conn, sid)                     # τήρηση 5ετίας


def test_document_html_contains_client_data_and_no_secrets(conn):
    from taxmatch.aml import cases, documents
    from taxmatch.business_profiles import credentials
    service.add(conn, AFM, "ΔΟΚΙΜΗ ΙΚΕ")
    conn.execute("UPDATE businesses SET legal_form='ΙΚΕ' WHERE afm=?", (AFM,))
    credentials.set_(conn, AFM, "user1", "SuperSecret!9")
    store.save_profile(conn, AFM, pep_status="domestic", purpose="Τήρηση βιβλίων")
    aid = store.save_assessment(conn, AFM, factors=["a_simple"], overrides=["o_pep"], approved_by="Διαχειριστής")
    cid = cases.save(conn, {"kind": "suspicion", "received_on": "2026-09-01", "description": "έμβασμα", "afm": AFM})
    pages = [documents.declaration_html(conn, AFM), documents.questionnaire_html(conn, AFM),
             documents.assessment_html(conn, aid), documents.engagement_clauses_html(conn, AFM),
             documents.case_file_html(conn, cid)]
    for h in pages:
        assert "ΔΟΚΙΜΗ ΙΚΕ" in h and "SuperSecret" not in h and "user1" not in h
    assert "Πραγματικοί δικαιούχοι" in pages[0]                   # νομικό πρόσωπο -> πίνακας δικαιούχων
    assert "ΥΨΗΛΟΣ" in pages[2] and "άρθρο 27" in pages[4] and "ΑΑΔΕ" not in pages[4].split("Αρχή")[0][-50:]


# ------------------------------------------------------------------ δικαιούχοι / πρότυπα Word

def test_ubos_carry_their_risk_to_the_client():
    f, o = model.factors_from_ubos([{"name": "Α", "pep": "no", "country_risk": "gr"},
                                    {"name": "Β", "pep": "family", "country_risk": "weak"}])
    assert "o_pep" in o and "Β" in o["o_pep"] and f == {"b_weak": "πραγματικός δικαιούχος Β"}
    _f, o = model.factors_from_ubos([{"name": "Γ", "country_risk": "sanctions"}])
    assert "o_sanctions" in o


def test_profile_keeps_rep_and_clean_ubos(conn):
    store.save_profile(conn, AFM, pep_status="no", rep={"name": "Εκπρόσωπος", "afm": "1", "bogus": "x"},
                       ubos=[{"name": "Δικαιούχος", "percent": "60", "pep": "nope", "country_risk": "mars"}, {"name": ""}])
    p = store.get_profile(conn, AFM)
    assert p["rep"] == {"name": "Εκπρόσωπος", "afm": "1"}
    assert len(p["ubos"]) == 1 and p["ubos"][0]["pep"] == "" and p["ubos"][0]["country_risk"] == ""


def _docx_text(data: bytes) -> str:
    import html, io, re, zipfile
    x = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()
    return html.unescape(re.sub(r"<[^>]+>", " ", x))


def test_render_handles_fields_split_across_runs_and_taxis_aliases():
    import io, zipfile
    from taxmatch.aml import templating
    xml = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
           '<w:p><w:r><w:t>Πελάτης: ${cli</w:t></w:r><w:r><w:rPr><w:b/></w:rPr><w:t>ent_name}</w:t></w:r></w:p>'
           '<w:p><w:r><w:t>ΑΦΜ ${mafm_etairias} {{ today }} ${agnosto} &amp; τέλος</w:t></w:r></w:p>'
           '<w:p><w:r><w:t>Χωρίς πεδία</w:t></w:r></w:p></w:body></w:document>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", xml)
    out, unknown = templating.render(buf.getvalue(), {"client_name": "Α & Β ΙΚΕ", "client_afm": "123", "today": "01/01/2026"})
    text = _docx_text(out)
    assert "Α & Β ΙΚΕ" in text and "ΑΦΜ 123 01/01/2026" in text and "& τέλος" in text and "Χωρίς πεδία" in text
    assert "${" not in text and unknown == ["agnosto"]


def test_default_templates_have_only_known_fields_and_generate_by_audience(conn):
    from taxmatch.aml import templating
    service.add(conn, AFM, "ΔΟΚΙΜΗ ΙΚΕ")
    conn.execute("UPDATE businesses SET legal_form='ΙΚΕ' WHERE afm=?", (AFM,))
    settings_store.set_value(conn, "aml_office_info", "ΓΡΑΦΕΙΟ Χ\nΟΔΟΣ 1")
    store.save_profile(conn, AFM, pep_status="no", rep={"name": "ΕΚΠΡΟΣΩΠΟΣ Α"},
                       ubos=[{"name": "ΔΙΚΑΙΟΥΧΟΣ Β", "percent": "100"}])
    store.save_assessment(conn, AFM, factors=["a_simple", "b_eu"])
    fields = templating.build_fields(conn, AFM)
    for slot in templating.SLOTS:
        d, a = slot.split(":")
        _out, unknown = templating.render(templating.default_template(d, a), fields)
        assert unknown == [], slot
    data, unknown, slot = templating.generate(conn, "declaration", AFM)
    text = _docx_text(data)
    assert slot == "declaration:legal" and not unknown
    assert "ΔΟΚΙΜΗ ΙΚΕ" in text and "ΕΚΠΡΟΣΩΠΟΣ Α" in text and "ΔΙΚΑΙΟΥΧΟΣ Β" in text and "ΓΡΑΦΕΙΟ Χ" in text
    data, _u, slot = templating.generate(conn, "assessment", AFM)
    assert slot == "assessment:legal" and "ΧΑΜΗΛΟΣ" in _docx_text(data)


def test_uploaded_template_replaces_default_and_can_be_reset(conn):
    from taxmatch.aml import templating
    with pytest.raises(ValueError):
        templating.save_template(conn, "declaration:natural", "x.docx", b"not a zip")
    custom = templating.default_template("agreement", "natural")
    templating.save_template(conn, "declaration:natural", "Δική μου δήλωση.docx", custom)
    data, name, is_custom = templating.get_template(conn, "declaration:natural")
    assert is_custom and name == "Δική μου δήλωση.docx" and data == custom
    templating.reset_template(conn, "declaration:natural")
    assert templating.get_template(conn, "declaration:natural")[2] is False
