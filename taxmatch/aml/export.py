"""Εξαγωγή σε Excel: φύλλο αξιολόγησης πελάτη (υπογράφεται και μπαίνει στον ατομικό φάκελο) και συγκεντρωτική
κατάσταση κατάταξης πελατών (άρθρο 35 παρ. 2) — με αριθμό έκδοσης/ημερομηνία, όπως ζητά ο έλεγχος."""
from __future__ import annotations

import sqlite3
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import settings_store
from . import model, store
from .content import (ASSESSMENT_LABEL, FEE_PAYMENTS, FILE_STATUS_LABEL, KYC_ITEMS, LAW_VERSION, MEASURES, TX_LABELS,
                      UBO_LABEL, documents_for, required_documents)

_BOLD = Font(bold=True)
_TITLE = Font(bold=True, size=14)
_HEAD_FILL = PatternFill("solid", fgColor="DDE7F0")
_CAT_FILL = {model.LOW: "D9F2E3", model.MEDIUM: "FFF1C9", model.HIGH: "F8D7DA"}


def _gr(iso: str) -> str:
    try:
        return date.fromisoformat(iso[:10]).strftime("%d/%m/%Y")
    except (TypeError, ValueError):
        return iso or ""


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def assessment_sheet(conn: sqlite3.Connection, assessment_id: int, path: Path) -> Path:
    a = store.get_assessment(conn, assessment_id)
    if not a:
        raise KeyError(assessment_id)
    profile = store.get_profile(conn, a["afm"])
    res = model.compute(a["factors"], a["overrides"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Αξιολόγηση"
    _widths(ws, [70, 26, 12, 12, 12])
    wrap = Alignment(wrap_text=True, vertical="top")

    ws.append(["ΦΥΛΛΟ ΑΞΙΟΛΟΓΗΣΗΣ ΚΙΝΔΥΝΟΥ ΠΕΛΑΤΗ — ΔΕΟΥΣΑ ΕΠΙΜΕΛΕΙΑ"])
    ws["A1"].font = _TITLE
    ws.append([LAW_VERSION])
    ws.append([])
    for label, value in [
        ("Πελάτης", a["client_name"]), ("ΑΦΜ", a["afm"]),
        ("Είδος αξιολόγησης", ASSESSMENT_LABEL.get(a["kind"], a["kind"])),
        ("Ημερομηνία αξιολόγησης", _gr(a["assessed_on"])),
        ("Μοντέλο που εφαρμόστηκε", f"Μοντέλο {a['model']}"),
        ("ΠΕΠ", model.PEP_LABEL.get(profile["pep_status"], profile["pep_status"])),
        ("Πραγματικοί δικαιούχοι", " — ".join(x for x in (UBO_LABEL.get(profile["ubo_state"], ""), profile["ubo_notes"]) if x)),
        ("Σκοπός και φύση της σχέσης", profile["purpose"]),
        ("Εξόφληση αμοιβής γραφείου", dict(FEE_PAYMENTS).get(profile["fee_payment"], "")),
        ("Κατάσταση φακέλου", FILE_STATUS_LABEL.get(profile["file_status"], "")),
    ]:
        ws.append([label, value])
        ws.cell(ws.max_row, 1).font = _BOLD
    ws.append([])

    if profile["transactions"]:
        ws.append(["ΣΥΝΑΛΛΑΓΕΣ / ΠΡΟΕΛΕΥΣΗ ΚΕΦΑΛΑΙΩΝ"])
        ws.cell(ws.max_row, 1).font = _BOLD
        for i, tx in enumerate(profile["transactions"], 1):
            parts = [TX_LABELS[f].get(tx.get(f, ""), "") for f in ("amount", "channel", "frequency", "activity", "country")]
            ws.append([f"{i}. " + " · ".join(x for x in parts if x), tx.get("justification", "")])
            ws.cell(ws.max_row, 1).alignment = wrap
        ws.append([])

    ws.append(["Παράγοντας", "Παράρτημα / διάταξη", "Μοντέλο Α", "Μοντέλο Β", "Συντρέχει"])
    for c in ws[ws.max_row]:
        c.font, c.fill = _BOLD, _HEAD_FILL
    chosen = set(a["factors"])
    for axis, label in model.AXES:
        ws.append([label])
        ws.cell(ws.max_row, 1).font = _BOLD
        for f in model.FACTORS:
            if f.axis != axis:
                continue
            ws.append([f.text, f.ref, f.a, "—" if f.b is None else f.b, "ΝΑΙ" if f.key in chosen else "ΟΧΙ"])
            ws.cell(ws.max_row, 1).alignment = wrap
        ws.append([f"Βαθμός άξονα (Μοντέλο Β, 0–25)", "", "", res.subs_b.get(axis, 0), ""])
    ws.append([])
    ws.append(["Περιπτώσεις «ΥΠΕΡΙΣΧΥΕΙ» — μία αρκεί για ΥΨΗΛΟ κίνδυνο"])
    ws.cell(ws.max_row, 1).font = _BOLD
    for o in model.OVERRIDES:
        ws.append([o.text, o.ref, "", "", "ΝΑΙ" if o.key in a["overrides"] else "ΟΧΙ"])
        ws.cell(ws.max_row, 1).alignment = wrap
    ws.append([])

    ws.append(["ΑΠΟΤΕΛΕΣΜΑ"])
    ws.cell(ws.max_row, 1).font = _BOLD
    rows = [
        ("Μοντέλο Α — άθροισμα / κατηγορία", f"{res.sum_a:+d}", model.CATEGORY_LABEL[res.cat_a]),
        ("Μοντέλο Β — σύνολο 0–100 / κατηγορία", str(res.total_b), model.CATEGORY_LABEL[res.cat_b]),
        ("Κατάταξη μοντέλου", model.CATEGORY_LABEL[a["model_category"]], ""),
        ("ΤΕΛΙΚΗ ΚΑΤΑΤΑΞΗ", model.CATEGORY_LABEL[a["final_category"]], model.DD_LABEL[a["final_category"]]),
    ]
    for r in rows:
        ws.append(list(r))
    ws.cell(ws.max_row, 1).font = _BOLD
    ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=_CAT_FILL[a["final_category"]])
    if a["escalation_note"]:
        ws.append(["Αιτιολογία αύξησης κατάταξης", a["escalation_note"]])
    ws.append(["Μέτρα", MEASURES[a["final_category"]]])
    ws.cell(ws.max_row, 2).alignment = wrap
    for note in res.notes:
        ws.append(["Σημείωση", note])
        ws.cell(ws.max_row, 2).alignment = wrap
    ws.append(["Αιτιολογία κατάταξης", a["justification"]])
    ws.cell(ws.max_row, 2).alignment = wrap
    ws.append(["Επόμενη επανεξέταση", _gr(a["next_review"])])
    ws.append([])
    ws.append(["ΕΓΓΡΑΦΑ ΦΑΚΕΛΟΥ", "Παραλήφθηκε"])
    ws.cell(ws.max_row, 1).font = _BOLD
    b = conn.execute("SELECT legal_form, activity_state FROM businesses WHERE afm=?", (a["afm"],)).fetchone()
    legal_form, activity = (b["legal_form"], b["activity_state"]) if b else ("", "")
    kind = store.effective_kind(profile, legal_form, activity)
    required = set(required_documents(legal_form, kind))
    for key, text in documents_for(legal_form, kind):
        ws.append([text + ("" if key in required else " (προαιρετικό)"), _gr(profile["docs"].get(key, "")) or "—"])
    ws.append([])
    ws.append(["ΛΙΣΤΑ KYC", "Ολοκληρώθηκε"])
    ws.cell(ws.max_row, 1).font = _BOLD
    for key, text, ref in KYC_ITEMS:
        ws.append([f"{text} ({ref})" if ref else text, _gr(profile["kyc"].get(key, "")) or "—"])
    ws.append([])
    ws.append(["Αξιολογητής", a["assessor"] or settings_store.get(conn, "aml_officer")])
    ws.append(["Έγκριση ανώτερου στελέχους (υψηλός κίνδυνος / ΠΕΠ)", a["approved_by"]])
    ws.append(["Υπογραφή", ""])
    ws.append([])
    ws.append(["Το φύλλο υπογράφεται, χρονολογείται και τίθεται στον ατομικό φάκελο του πελάτη."])
    path = Path(path)
    wb.save(path)
    return path


def summary_sheet(conn: sqlite3.Connection, path: Path, version: str = "") -> Path:
    """Συγκεντρωτική κατάσταση κατάταξης: μία γραμμή ανά πελάτη (και πρώην πελάτη με ιστορικό)."""
    rows = store.overview(conn)
    wb = Workbook()
    ws = wb.active
    ws.title = "Κατάσταση πελατών"
    _widths(ws, [12, 40, 14, 14, 16, 26, 14, 16, 8, 10, 40, 14])
    now = datetime.now()
    ws.append(["ΣΥΓΚΕΝΤΡΩΤΙΚΗ ΚΑΤΑΣΤΑΣΗ ΚΑΤΑΤΑΞΗΣ ΠΕΛΑΤΩΝ (άρθρο 35 παρ. 2 ν. 4557/2018)"])
    ws["A1"].font = _TITLE
    ws.append([f"Έκδοση {version or now.strftime('%Y.%m.%d')} — {now.strftime('%d/%m/%Y')} — Μοντέλο "
               f"{store.office_model(conn)} — Υπεύθυνος συμμόρφωσης: {settings_store.get(conn, 'aml_officer') or '—'}"])
    ws.append([])
    head = ["ΑΦΜ", "Επωνυμία", "Κατηγορία", "Βαθμολογία", "Δέουσα επιμέλεια", "ΠΕΠ", "Τελ. αξιολόγηση",
            "Επόμενη επανεξέταση", "KYC", "Έγγραφα", "Κατάσταση φακέλου", "Σχέση"]
    ws.append(head)
    for c in ws[ws.max_row]:
        c.font, c.fill = _BOLD, _HEAD_FILL
    status_label = {"missing": "ΧΩΡΙΣ ΑΞΙΟΛΟΓΗΣΗ", "overdue": "ΕΚΠΡΟΘΕΣΜΗ", "due_soon": "", "ok": ""}
    for r in rows:
        a: dict[str, Any] | None = r["assessment"]
        score = ""
        if a:
            score = f"{a['total_b']}/100" if a["model"] == "B" else f"{a['sum_a']:+d}"
        ws.append([
            r["afm"], r["name"], model.CATEGORY_LABEL.get(r["category"], status_label["missing"]), score,
            model.DD_SHORT.get(r["category"], ""), model.PEP_LABEL.get(r["pep_status"], ""),
            _gr(a["assessed_on"]) if a else "",
            (_gr(r["next_review"]) + (" — " + status_label[r["status"]] if status_label.get(r["status"]) else "")) if a else "",
            f"{r['kyc_done']}/{r['kyc_total']}", f"{r['docs_done']}/{r['docs_total']}",
            FILE_STATUS_LABEL.get(r["file_status"], ""), "ενεργή" if r["is_client"] else "λήξη (τήρηση 5ετίας)",
        ])
        if r["category"]:
            ws.cell(ws.max_row, 3).fill = PatternFill("solid", fgColor=_CAT_FILL[r["category"]])
    ws.freeze_panes = "A5"
    path = Path(path)
    wb.save(path)
    return path
