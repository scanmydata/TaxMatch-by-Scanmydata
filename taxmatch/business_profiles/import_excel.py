"""Μαζική εισαγωγή πελατών από Excel/CSV: ΑΦΜ + προαιρετικά επωνυμία, ΚΑΔ, **κωδικούς TAXISnet** και **Γ.Ε.ΜΗ.**

Αναγνωρίζονται δύο μορφές (όπως στο timologio downloader του ScanMyData):
* πίνακας με επικεφαλίδες (μία γραμμή ανά πελάτη)·
* μπλοκ ανά υπόχρεο — η εκτύπωση «Κωδικοί Υπηρεσιών μέσω Internet» των Hyper/Extra/TaxSystem, με γραμμή «Taxis Net»
  (και «Γ.Ε.ΜΗ.») (χρήστης, συνθηματικό)· οι τιμές διαβάζονται ΘΕΣΙΑΚΑ ανάμεσα στα μη κενά κελιά, όχι από σταθερές στήλες.

Το ταίριασμα επικεφαλίδων είναι whole-string (ποτέ substring): μια στήλη «Κλειδί myDATA» δεν πρέπει ποτέ να
διαβαστεί ως κωδικός TAXISnet. Διαβάζονται ΜΟΝΟ οι κωδικοί TAXISnet, Γ.Ε.ΜΗ. (σύνδεση στο businessportal για τη
Δέουσα επιμέλεια), myDATA REST API (όνομα χρήστη + κλειδί «Api myData», όπως στο Timologio Downloader — ΠΟΤΕ το
«Συνθηματικό myData», που είναι ο κωδικός της ιστοσελίδας, ούτε το κλειδί του e-timologio, άλλο προϊόν) και — ρητό
αίτημα χρήστη 2026-10-22, για τον έλεγχο ΑΠΔ — κωδικοί **ΙΚΑ εργοδότη** + **ΑΜΕ**· όσοι άλλων υπηρεσιών (ΙΚΑ
ασφαλισμένου, ΣΕΠΕ, ΟΑΕΔ, Κτηματολόγιο…) δεν διαβάζονται ποτέ.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from typing import Iterable, Optional

from ..identifiers import is_valid_afm, normalize_afm
from ..textutil import strip_accents

FIELD_ALIASES: dict[str, set[str]] = {
    "afm": {"αφμ", "α φ μ", "afm", "vat", "vat number", "vatnumber", "tax id", "taxid", "αφμ υποχρεου"},
    "name": {"επωνυμια επωνυμο", "επωνυμια", "επωνυμο", "ονομασια", "name", "company", "πελατης", "επιχειρηση",
             "ονομα επιχειρησης", "onomasia"},
    # «Όνομα» (μικρό όνομα) — στο Excel των λογιστικών προγραμμάτων το «Επωνυμία/Επώνυμο» έχει ΜΟΝΟ το επώνυμο των
    # φυσικών προσώπων και το μικρό όνομα είναι χωριστή στήλη· χωρίς αυτήν οι πελάτες εισάγονταν ως «ΑΓΓΕΛΙΚΑΚΗΣ».
    "first_name": {"ονομα", "μικρο ονομα", "first name", "firstname", "given name"},
    "kad": {"καδ", "κ α δ", "kad", "καδ κυριο", "κυριος καδ", "καδ κυριος"},
    "taxis_user": {"ονομα χρηστη taxisnet", "χρηστης taxisnet", "taxisnet user", "taxisnet username", "taxis user",
                   "ονομα χρηστη taxis", "username taxisnet", "taxis net user", "ονομα χρηστη taxis net",
                   "taxisnet ονομα χρηστη", "χρηστης taxis"},
    "taxis_pass": {"συνθηματικο taxisnet", "κωδικος taxisnet", "taxisnet password", "taxis password",
                   "κωδικος προσβασης taxisnet", "συνθηματικο taxis", "κωδικος taxis", "taxis net password",
                   "κωδικος προσβασης taxis net", "taxisnet κωδικος"},
    # «Όνομα χρήστη Γ.Ε.Μ.Η.» / «Συνθηματικό χρήστη Γ.Ε.Μ.Η.» της εξαγωγής «Κωδικοί Υπόχρεων» (TaxSystem) — επαληθεύτηκε
    # σε πραγματικό αρχείο γραφείου (2026-10-22, μόνο οι επικεφαλίδες).
    "gemi_user": {"ονομα χρηστη γ ε μ η", "ονομα χρηστη γεμη", "χρηστης γ ε μ η", "χρηστης γεμη", "gemi user",
                  "gemi username", "username γεμη", "ονομα χρηστη γ ε μη", "γ ε μ η ονομα χρηστη"},
    "gemi_pass": {"συνθηματικο χρηστη γ ε μ η", "συνθηματικο γ ε μ η", "συνθηματικο χρηστη γεμη", "συνθηματικο γεμη",
                  "κωδικος γ ε μ η", "κωδικος γεμη", "gemi password", "κωδικος προσβασης γεμη", "συνθηματικο γ ε μη",
                  "γ ε μ η συνθηματικο"},
    # Σημαίες για τις προθεσμίες (obligations.liability): «Μισθοδοσία» (True/False) του TaxSystem = εργοδότης (ΑΠΔ/Ε4)·
    # στοιχεία σύνδεσης Intrastat = υποβάλλει Intrastat. ΔΕΝ κρατάμε τις τιμές των κωδικών Intrastat — μόνο ότι υπάρχουν.
    "employer": {"μισθοδοσια", "εργοδοτης", "απασχολει προσωπικο", "employer", "payroll"},
    # myDATA REST API — ίδια aliases με το timologio (excel/aliases.py). «Subscription key e-timologio» ΔΕΝ είναι εδώ:
    # το παλιό substring ταίριασμα το άρπαζε ως κλειδί myDATA → 403 σε κάθε πελάτη.
    "mydata_user": {"ονομα χρηστη mydata", "ονομα χρηστη my data", "aade user id", "aade user", "χρηστης mydata"},
    "mydata_key": {"api mydata", "api my data", "ocp apim subscription key", "subscription key mydata",
                   "κλειδι api mydata"},
    "intrastat_user": {"email intrastat", "ονομα χρηστη intrastat", "χρηστης intrastat"},
    # Κωδικοί ΙΚΑ ΕΡΓΟΔΟΤΗ (όχι ασφαλισμένου): «Όνομα χρήστη/Συνθηματικό εργοδότη Ι.Κ.Α.» της εξαγωγής «Κωδικοί Υπόχρεων»
    # του TaxSystem (επαληθεύτηκε σε πραγματικό αρχείο 2026-10-22, μόνο οι επικεφαλίδες)· ίδια με το ScanMyData Tax
    # Center (`ColumnAliases.kt`). ΑΜΕ: δεν υπάρχει σε εκείνο το αρχείο — δεκτός αν κάποιο τον έχει.
    "ika_user": {"ονομα χρηστη εργοδοτη ι κ α", "ονομα χρηστη εργοδοτη ικα", "χρηστης εργοδοτη ικα",
                 "ονομα χρηστη ικα εργοδοτη", "ονομα χρηστη ι κ α εργοδοτη", "ika employer user"},
    "ika_pass": {"συνθηματικο εργοδοτη ι κ α", "συνθηματικο εργοδοτη ικα", "κωδικος εργοδοτη ικα",
                 "συνθηματικο χρηστη εργοδοτη ι κ α", "συνθηματικο ικα εργοδοτη", "ika employer password"},
    "ame": {"αμε", "α μ ε", "αριθμος μητρωου εργοδοτη", "αρ μητρωου εργοδοτη", "αμ εργοδοτη", "αμε εργοδοτη",
            "α μ εργοδοτη"},
    # ΑΜΚΑ (ΚΕΑΟ / τρέχουσες εισφορές e-ΕΦΚΑ φυσικών προσώπων — Tax Center ColumnAliases)
    "amka": {"αμκα", "α μ κ α", "amka"},
    # Μέλη εταιρειών (members.py): στη γραμμή του μέλους ο ΑΦΜ της εταιρείας του, ή στη γραμμή της εταιρείας οι ΑΦΜ
    # των μελών (λίστα). ΔΕΝ έχει επαληθευτεί με πραγματικό αρχείο — αν ένα πρόγραμμα το γράφει αλλιώς, alias εδώ.
    "member_of": {"αφμ εταιρειας", "εταιρεια αφμ", "αφμ νομικου προσωπου", "μελος της εταιρειας", "μελος εταιρειας",
                  "αφμ εταιρειας μελους"},
    "member_afms": {"μελη", "αφμ μελων", "εταιροι", "εταιροι αφμ", "μετοχοι", "μελη αφμ", "διαχειριστες"},
}
#: Γενικές στήλες: δεκτές ΜΟΝΟ σε απλό αρχείο χωρίς ρητές στήλες TAXISnet (ώστε να δουλεύει το πρότυπό μας
#: «ΑΦΜ | Επωνυμία | Όνομα χρήστη | Κωδικός»), όχι μέσα σε πλατύ πίνακα με δεκάδες κωδικούς διαφορετικών υπηρεσιών.
GENERIC_USER = {"ονομα χρηστη", "χρηστης", "username", "user"}
GENERIC_PASS = {"κωδικος", "κωδικος προσβασης", "συνθηματικο", "password"}
SIMPLE_SHEET_MAX_COLS = 8

BLOCK = re.compile(r"^(\d{3,6})\s+(.+?)\s+(\d{9})$")
_TAXIS_LABELS = ("taxis net", "taxisnet", "taxis")
_GEMI_LABELS = ("γ ε μ η", "γ ε μη", "γεμη", "gemi")
_IKA_LABELS = ("ι κ α", "ικα", "εφκα")
_BLOCK_SPAN = 40                    # μέχρι το επόμενο μπλοκ· στο Hyper/Extra ένα μπλοκ φτάνει τις 15 γραμμές
_HEADER_SCAN = 15

TEMPLATE_HEADERS = ["ΑΦΜ", "Επωνυμία", "ΚΑΔ", "Όνομα χρήστη TAXISnet", "Κωδικός TAXISnet", "Όνομα χρήστη Γ.Ε.Μ.Η.",
                    "Συνθηματικό χρήστη Γ.Ε.Μ.Η."]


def norm_header(v: object) -> str:
    """Ελληνικά/λατινικά, χωρίς τόνους/σημεία στίξης, μικρά, μονά κενά: 'Α.Φ.Μ.' -> 'α φ μ'."""
    s = strip_accents(str(v or "")).lower()
    s = re.sub(r"[^0-9a-zα-ω]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


_LOOKUP = {alias: f for f, aliases in FIELD_ALIASES.items() for alias in aliases}


@dataclass
class ImportRow:
    afm: str
    name: str = ""
    kads: list[str] = field(default_factory=list)
    checksum_ok: bool = True
    taxis_user: str = ""
    taxis_pass: str = ""
    gemi_user: str = ""
    gemi_pass: str = ""
    mydata_user: str = ""
    mydata_key: str = ""
    ika_user: str = ""
    ika_pass: str = ""
    ame: str = ""
    amka: str = ""
    member_of: list[str] = field(default_factory=list)
    member_afms: list[str] = field(default_factory=list)
    employer: Optional[bool] = None                  # None = άγνωστο (δεν υπάρχει στήλη / κενό κελί)
    intrastat: Optional[bool] = None                 # True μόνο όταν υπάρχουν στοιχεία Intrastat· ποτέ False από απουσία

    @property
    def has_credentials(self) -> bool:
        return bool(self.taxis_user and self.taxis_pass)

    @property
    def has_gemi(self) -> bool:
        return bool(self.gemi_user and self.gemi_pass)

    @property
    def has_mydata(self) -> bool:
        return bool(self.mydata_user and self.mydata_key)

    @property
    def has_ika(self) -> bool:
        return bool(self.ika_user and self.ika_pass)


@dataclass
class ImportResult:
    rows: list[ImportRow] = field(default_factory=list)
    invalid: list[tuple[int, str]] = field(default_factory=list)      # (γραμμή αρχείου, τιμή)
    duplicates: int = 0
    warnings: list[str] = field(default_factory=list)
    detected: str = ""                                                # 'table' | 'blocks'

    @property
    def credential_count(self) -> int:
        return sum(1 for r in self.rows if r.has_credentials)

    @property
    def gemi_count(self) -> int:
        return sum(1 for r in self.rows if r.has_gemi)

    @property
    def mydata_count(self) -> int:
        return sum(1 for r in self.rows if r.has_mydata)

    @property
    def ika_count(self) -> int:
        return sum(1 for r in self.rows if r.has_ika or r.ame)


def full_name(surname: str, first_name: str) -> str:
    """«ΕΠΩΝΥΜΟ ΟΝΟΜΑ» για φυσικά πρόσωπα· οι εταιρείες (κενό μικρό όνομα) μένουν ως έχουν."""
    surname, first_name = " ".join(surname.split()), " ".join(first_name.split())
    if not first_name or strip_accents(surname).upper().endswith(strip_accents(first_name).upper()):
        return surname
    return f"{surname} {first_name}".strip()


def _flag(v: object) -> Optional[bool]:
    """True/False, ΝΑΙ/ΟΧΙ, 1/0, Χ — κενό = άγνωστο."""
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return v
    s = strip_accents(str(v)).strip().lower()
    if s in ("true", "ναι", "nai", "yes", "y", "1", "x", "χ", "✓"):
        return True
    if s in ("false", "οχι", "oxi", "no", "n", "0", "-"):
        return False
    return None


_MYDATA_KEY_SHAPE = re.compile(r"^[0-9a-fA-F]{32}$")
#: Στήλη που ΚΑΝΟΝΙΚΑ έχει τον κωδικό της ιστοσελίδας myDATA — αλλά κάποιοι λογιστές βάζουν εκεί κατά λάθος το κλειδί
#: API. Διαβάζεται ΜΟΝΟ αν λείπει η «Api myData» ΚΑΙ η τιμή είναι αδιαμφισβήτητα κλειδί (32 hex) — ίδιος κανόνας με το
#: timologio (`excel/aliases.SECONDARY_KEY_HEADERS`).
_SECONDARY_KEY_HEADERS = {"συνθηματικο mydata", "συνθηματικο my data"}


def _mydata_key(row: tuple, mapping: dict[str, int], heads: list[str]) -> str:
    idx = mapping.get("mydata_key")
    value = _cell(row[idx]) if idx is not None and idx < len(row) else ""
    if value:
        return value
    for j, h in enumerate(heads):
        if h in _SECONDARY_KEY_HEADERS and j < len(row) and _MYDATA_KEY_SHAPE.match(_cell(row[j])):
            return _cell(row[j])
    return ""


def _afm_list(v: object) -> list[str]:
    out = []
    for part in re.split(r"[,;/\n\s]+", str(v or "")):
        afm = normalize_afm(part)
        if afm and afm not in out:
            out.append(afm)
    return out


def _split_kads(v: object) -> list[str]:
    return [p.strip() for p in re.split(r"[;,/\n]+", str(v or "")) if re.search(r"\d{2}", p)]


def _cell(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _find_header(rows: list[tuple]) -> Optional[tuple[int, dict[str, int]]]:
    for i, row in enumerate(rows[:_HEADER_SCAN]):
        heads = [norm_header(c) for c in row]
        mapping: dict[str, int] = {}
        for j, h in enumerate(heads):
            f = _LOOKUP.get(h)
            if f and f not in mapping:
                mapping[f] = j
        if "afm" not in mapping:
            continue
        if not ("taxis_user" in mapping or "taxis_pass" in mapping) and len([h for h in heads if h]) <= SIMPLE_SHEET_MAX_COLS:
            for j, h in enumerate(heads):
                if h in GENERIC_USER and "taxis_user" not in mapping:
                    mapping["taxis_user"] = j
                elif h in GENERIC_PASS and "taxis_pass" not in mapping:
                    mapping["taxis_pass"] = j
        return i, mapping
    return None


def _parse_blocks(rows: list[tuple]) -> Optional[ImportResult]:
    """Μπλοκ «κωδικός επωνυμία ΑΦΜ» + γραμμή «Taxis Net» με χρήστη/συνθηματικό."""
    texts = [" ".join(_cell(c) for c in row if _cell(c)) for row in rows]
    if not any(BLOCK.match(t) for t in texts[:120]):
        return None
    res = ImportResult(detected="blocks")
    seen: set[str] = set()
    i = 0
    while i < len(rows):
        m = BLOCK.match(texts[i])
        if not m:
            i += 1
            continue
        _code, name, raw_afm = m.groups()
        afm = normalize_afm(raw_afm)
        user = pwd = g_user = g_pwd = m_user = m_key = i_user = i_pwd = ame = ""
        intrastat: Optional[bool] = None
        j = i + 1
        while j < min(i + _BLOCK_SPAN, len(rows)):
            if j > i + 1 and BLOCK.match(texts[j]):
                break
            low = strip_accents(texts[j]).lower()
            if any(lbl in low for lbl in _TAXIS_LABELS) and "mydata" not in low and "βιβλια" not in low:
                vals = [_cell(c) for c in rows[j] if _cell(c)]
                if len(vals) >= 3:
                    user, pwd = vals[1], vals[2]
            elif "mydata" in low or "ηλεκτρονικα βιβλια" in low:
                vals = [_cell(c) for c in rows[j] if _cell(c)]
                if len(vals) >= 3:
                    m_user, m_key = vals[1], vals[2]          # «Ηλεκτρονικά Βιβλία Α.Α.Δ.Ε. (myDATA)»: χρήστης, κλειδί API
            elif "εργοδοτ" in norm_header(texts[j]) and any(lbl in norm_header(texts[j]) for lbl in _IKA_LABELS) \
                    and "ασφαλισμεν" not in low:
                vals = [_cell(c) for c in rows[j] if _cell(c)]
                if len(vals) >= 3:
                    i_user, i_pwd = vals[1], vals[2]              # «Ι.Κ.Α. Εργοδότη»: χρήστης, συνθηματικό
            elif re.match(r"^(α μ ε|αμε)( |$)", norm_header(texts[j])):
                vals = [_cell(c) for c in rows[j] if _cell(c)]
                if len(vals) >= 2 and re.fullmatch(r"\d{6,10}", vals[1]):
                    ame = vals[1]
            elif "intrastat" in low and len([c for c in rows[j] if _cell(c)]) >= 2:
                intrastat = True                         # γραμμή «Δηλώσεων Intrastat» με στοιχεία σύνδεσης
            elif any(lbl in norm_header(texts[j]) for lbl in _GEMI_LABELS):
                vals = [_cell(c) for c in rows[j] if _cell(c)]
                if len(vals) >= 3:
                    g_user, g_pwd = vals[1], vals[2]
            j += 1
        i = j
        if not afm:
            res.invalid.append((i, raw_afm))
            continue
        if afm in seen:
            # Ο ίδιος υπόχρεος σε δύο εταιρείες της εκτύπωσης: κρατάμε την πρώτη, αλλά συμπληρώνουμε κωδικούς που
            # έλειπαν από αυτήν (π.χ. Γ.Ε.ΜΗ. μόνο στο δεύτερο μπλοκ).
            res.duplicates += 1
            first = next(r for r in res.rows if r.afm == afm)
            if not first.has_credentials and user and pwd:
                first.taxis_user, first.taxis_pass = user, pwd
            if not first.has_gemi and g_user and g_pwd:
                first.gemi_user, first.gemi_pass = g_user, g_pwd
            if not first.has_mydata and m_user and m_key:
                first.mydata_user, first.mydata_key = m_user, m_key
            if not first.has_ika and i_user and i_pwd:
                first.ika_user, first.ika_pass = i_user, i_pwd
            first.ame = first.ame or ame
            continue
        seen.add(afm)
        res.rows.append(ImportRow(afm=afm, name=" ".join(name.split()), checksum_ok=is_valid_afm(afm),
                                  taxis_user=user, taxis_pass=pwd, gemi_user=g_user, gemi_pass=g_pwd,
                                  intrastat=intrastat, mydata_user=m_user, mydata_key=m_key,
                                  ika_user=i_user, ika_pass=i_pwd, ame=ame))
    return res


def parse_rows(table: Iterable[tuple]) -> ImportResult:
    rows = [tuple(r) for r in table]
    found = _find_header(rows)
    if found is None:
        blocks = _parse_blocks(rows)
        if blocks is not None:
            _finish(blocks)
            return blocks
    result = ImportResult(detected="table")
    heads: list[str] = []
    if found:
        start, mapping = found[0] + 1, found[1]
        heads = [norm_header(c) for c in rows[found[0]]]
    else:
        start, mapping = 0, {"afm": 0, "name": 1}
        result.warnings.append("Δεν βρέθηκε επικεφαλίδα «ΑΦΜ» — χρησιμοποιήθηκε η πρώτη στήλη.")

    def get(row: tuple, key: str) -> str:
        idx = mapping.get(key)
        return _cell(row[idx]) if idx is not None and idx < len(row) else ""

    seen: set[str] = set()
    for offset, row in enumerate(rows[start:], start=start + 1):
        if not row or all(_cell(c) == "" for c in row):
            continue
        afm = normalize_afm(row[mapping["afm"]] if mapping["afm"] < len(row) else None)
        if not afm:
            result.invalid.append((offset, get(row, "afm")))
            continue
        if afm in seen:
            result.duplicates += 1
            continue
        seen.add(afm)
        result.rows.append(ImportRow(afm=afm, name=full_name(get(row, "name"), get(row, "first_name")), kads=_split_kads(get(row, "kad")),
                                     checksum_ok=is_valid_afm(afm),
                                     taxis_user=get(row, "taxis_user"), taxis_pass=get(row, "taxis_pass"),
                                     gemi_user=get(row, "gemi_user"), gemi_pass=get(row, "gemi_pass"),
                                     employer=_flag(row[mapping["employer"]]) if mapping.get("employer") is not None
                                     and mapping["employer"] < len(row) else None,
                                     intrastat=True if get(row, "intrastat_user") else None,
                                     mydata_user=get(row, "mydata_user"), mydata_key=_mydata_key(row, mapping, heads),
                                     ika_user=get(row, "ika_user"), ika_pass=get(row, "ika_pass"), ame=get(row, "ame"),
                                     amka=get(row, "amka"), member_of=_afm_list(get(row, "member_of")),
                                     member_afms=_afm_list(get(row, "member_afms"))))
    _finish(result)
    return result


def _finish(result: ImportResult) -> None:
    bad = sum(1 for r in result.rows if not r.checksum_ok)
    if bad:
        result.warnings.append(f"{bad} ΑΦΜ δεν περνούν τον έλεγχο ψηφίου ελέγχου — εισάγονται αλλά ελέγξτε τα.")
    half = [r.afm for r in result.rows if bool(r.taxis_user) != bool(r.taxis_pass)]
    if half:
        result.warnings.append(f"{len(half)} πελάτες έχουν μόνο χρήστη ή μόνο κωδικό TAXISnet — οι κωδικοί τους αγνοούνται.")
        for r in result.rows:
            if r.afm in half:
                r.taxis_user = r.taxis_pass = ""
    half = [r.afm for r in result.rows if bool(r.mydata_user) != bool(r.mydata_key)]
    if half:
        result.warnings.append(f"{len(half)} πελάτες έχουν μόνο χρήστη ή μόνο κλειδί myDATA — αγνοούνται.")
        for r in result.rows:
            if r.afm in half:
                r.mydata_user = r.mydata_key = ""
    half = [r.afm for r in result.rows if bool(r.ika_user) != bool(r.ika_pass)]
    if half:
        result.warnings.append(f"{len(half)} πελάτες έχουν μόνο χρήστη ή μόνο κωδικό ΙΚΑ εργοδότη — αγνοούνται.")
        for r in result.rows:
            if r.afm in half:
                r.ika_user = r.ika_pass = ""
    half = [r.afm for r in result.rows if bool(r.gemi_user) != bool(r.gemi_pass)]
    if half:
        result.warnings.append(f"{len(half)} πελάτες έχουν μόνο χρήστη ή μόνο κωδικό Γ.Ε.ΜΗ. — οι κωδικοί τους αγνοούνται.")
        for r in result.rows:
            if r.afm in half:
                r.gemi_user = r.gemi_pass = ""


def _col_index(ref: str) -> int:
    n = 0
    for ch in ref:
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def _raw_xlsx_rows(content: bytes) -> list[tuple]:
    """Πρώτο φύλλο ενός .xlsx ΧΩΡΙΣ openpyxl (μόνο κείμενο/αριθμοί, κανένα στυλ). Κενά κελιά -> None στη θέση τους,
    ώστε η θεσιακή ανάγνωση των μπλοκ να βλέπει το ίδιο ό,τι και με το openpyxl."""
    import html as _html
    import zipfile
    z = zipfile.ZipFile(io.BytesIO(content))
    names = z.namelist()
    shared: list[str] = []
    if "xl/sharedStrings.xml" in names:
        xml = z.read("xl/sharedStrings.xml").decode("utf-8")
        for si in re.findall(r"<si>(.*?)</si>", xml, re.S):
            shared.append(_html.unescape("".join(re.findall(r"<t(?: [^>]*)?>(.*?)</t>", si, re.S))))
    sheet = "xl/worksheets/sheet1.xml"
    try:                                            # το «πρώτο» φύλλο κατά το workbook.xml, όχι κατά όνομα αρχείου
        wb = z.read("xl/workbook.xml").decode("utf-8")
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        rid = re.search(r'<sheet [^>]*r:id="([^"]+)"', wb).group(1)
        target = re.search(rf'<Relationship [^>]*Id="{re.escape(rid)}"[^>]*Target="([^"]+)"', rels) \
            or re.search(rf'<Relationship [^>]*Target="([^"]+)"[^>]*Id="{re.escape(rid)}"', rels)
        sheet = "xl/" + target.group(1).lstrip("/").removeprefix("xl/")
    except (KeyError, AttributeError):
        pass
    xml = z.read(sheet).decode("utf-8")
    out: list[tuple] = []
    for row in re.findall(r"<row[^>]*?(?:/>|>(.*?)</row>)", xml, re.S):
        cells: dict[int, object] = {}
        for ref, attrs, body in re.findall(r'<c r="([A-Z]+)\d+"([^>]*?)(?:/>|>(.*?)</c>)', row or "", re.S):
            kind = re.search(r'\bt="(\w+)"', attrs)
            kind = kind.group(1) if kind else ""
            v = re.search(r"<v>(.*?)</v>", body or "", re.S)
            if kind == "s" and v:
                val: object = shared[int(v.group(1))]
            elif kind == "inlineStr":
                val = _html.unescape("".join(re.findall(r"<t(?: [^>]*)?>(.*?)</t>", body or "", re.S)))
            elif v:
                raw = _html.unescape(v.group(1))
                val = raw
                if kind not in ("str", "e", "b"):
                    try:
                        num = float(raw)
                        val = int(num) if num.is_integer() else num
                    except ValueError:
                        pass
            else:
                continue
            cells[_col_index(ref)] = val
        out.append(tuple(cells.get(i) for i in range(max(cells) + 1)) if cells else ())
    return out


def parse_file(filename: str, content: bytes) -> ImportResult:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:                           # noqa: BLE001 — ό,τι κι αν σκάσει στο openpyxl, δοκιμάζουμε ωμά
            # Η εξαγωγή «Κωδικοί Υπηρεσιών μέσω Internet» του Hyper/Extra έχει άκυρο styles.xml («Colors must be aRGB
            # hex values» — επαληθεύτηκε σε πραγματικό αρχείο, 2026-10-22) και το openpyxl αρνείται όλο το αρχείο.
            # Τα στυλ δεν μας χρειάζονται: διαβάζουμε απευθείας τα κελιά.
            return parse_rows(_raw_xlsx_rows(content))
        try:
            # Μόνο το πρώτο φύλλο: τα επόμενα (π.χ. «Sheet (2)») συχνά έχουν τους ίδιους ΑΦΜ χωρίς κωδικούς.
            return parse_rows(wb.worksheets[0].iter_rows(values_only=True))
        finally:
            wb.close()
    if name.endswith((".csv", ".txt")):
        text = None
        for enc in ("utf-8-sig", "cp1253", "latin-1"):
            try:
                text = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        text = text or ""
        try:
            dialect = csv.Sniffer().sniff(text[:2000], delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
        return parse_rows(csv.reader(io.StringIO(text), dialect))
    raise ValueError("Υποστηρίζονται αρχεία .xlsx και .csv")


def template_xlsx() -> bytes:
    """Κενό πρότυπο με τις στήλες που αναγνωρίζονται, για όσους δεν έχουν εξαγωγή από πρόγραμμα λογιστικής."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Πελάτες"
    ws.append(TEMPLATE_HEADERS)
    for col, width in zip("ABCDEFG", (14, 38, 18, 26, 26, 26, 26)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
